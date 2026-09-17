"""生成用户手册与开发者手册（.docx）。运行：python -m scripts.build_manuals"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x36, 0x5D)
INK = RGBColor(0x21, 0x25, 0x29)
MUTED = RGBColor(0x5A, 0x62, 0x6E)
AMBER = RGBColor(0x8A, 0x5A, 0x12)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = "1B365D"

OUT_DIR = Path(__file__).resolve().parent.parent / "docs" / "manuals"


def _east_asia(run, name: str = "微软雅黑") -> None:
    run.font.name = "Calibri"
    r = run._element.get_or_add_rPr()
    rFonts = r.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        r.append(rFonts)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), "Calibri")


def _shade(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _set_cell_border(cell, color: str = RULE) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _page_number_run(paragraph) -> None:
    run = paragraph.add_run()
    _east_asia(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _setup_doc(title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = INK
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    pf = styles["Normal"].paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)

    for name, size, space_before in (("Heading 1", 18, 18), ("Heading 2", 14, 14), ("Heading 3", 12, 10)):
        st = styles[name]
        st.font.color.rgb = NAVY
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.name = "微软雅黑"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run(f"{title}  ·  仅供个人研究辅助")
    _east_asia(hr)
    hr.font.size = Pt(8)
    hr.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("不构成投资建议  ·  第 ")
    _east_asia(fr)
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED
    _page_number_run(footer)
    fr2 = footer.add_run(" 页")
    _east_asia(fr2)
    fr2.font.size = Pt(8)
    fr2.font.color.rgb = MUTED
    return doc


def add_run(p, text, *, size=11, bold=False, color=INK, italic=False):
    run = p.add_run(text)
    _east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return run


def para(doc, text, *, size=11, bold=False, color=INK, align=None, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    if align:
        p.alignment = align
    add_run(p, text, size=size, bold=bold, color=color)
    return p


def bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        add_run(p, item, size=11)


def numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.clear()
        add_run(p, item, size=11)


def callout(doc, title: str, body: str, fill: str = "FFF7E6") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _shade(cell, fill)
    _set_cell_border(cell, "D4A017" if fill == "FFF7E6" else RULE)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(4)
    add_run(p1, title, size=11, bold=True, color=AMBER if fill == "FFF7E6" else NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    add_run(p2, body, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def cover(doc, kicker: str, title: str, subtitle: str, meta: list[str]) -> None:
    banner = doc.add_table(rows=1, cols=1)
    cell = banner.cell(0, 0)
    _shade(cell, "1B365D")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_run(p, kicker, size=11, bold=True, color=WHITE)
    p2 = cell.add_paragraph()
    add_run(p2, title, size=26, bold=True, color=WHITE)
    p3 = cell.add_paragraph()
    add_run(p3, subtitle, size=12, color=RGBColor(0xD6, 0xDE, 0xE8))
    doc.add_paragraph()
    for line in meta:
        para(doc, line, size=11, color=MUTED, space_after=2)
    doc.add_paragraph()
    callout(
        doc,
        "使用前请先读这一段",
        "本系统全部输出仅供个人研究与学习，不构成投资建议、不承诺收益、不代客理财、不自动下单。"
        "最终交易决策及盈亏由使用者自行承担。建议卡片上的动作（减仓/止损/观察等）是规则引擎给出的信息性提示，不是下单指令。",
        fill="FEF2F2",
    )


def make_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        _shade(cell, "1B365D")
        _set_cell_border(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        add_run(p, h, size=10, bold=True, color=WHITE)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            _set_cell_border(cell, "D6DEE8")
            if r_i % 2 == 1:
                _shade(cell, "F5F7FA")
            cell.text = ""
            p = cell.paragraphs[0]
            add_run(p, val, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def build_user() -> Document:
    doc = _setup_doc("用户使用说明书")
    cover(
        doc,
        "个人量化研究系统  ·  stock-quant-system",
        "用户使用说明书",
        "持仓驾驶舱 · 建议卡片 · 掘金扫描 · 行情图表 · 风险 · AI 解释 · 复盘 · 名词解释",
        [
            "版本 1.1    编写日期 2026-08-26（第二轮可用性测试配套）",
            "适用对象：自己看盘、自己下单的个人使用者",
            "配套：Phase 0–4 已验收的本地 Streamlit 应用 + 一键启动器",
            "人工测试步骤见同目录《人工可用性测试指南》",
        ],
    )

    doc.add_heading("1. 这份说明书解决什么问题", level=1)
    para(
        doc,
        "告诉你：怎么打开系统、每个页面干什么、建议卡片上的词是什么意思、哪些数字可以当真、哪些绝对不能当真。"
        "它不教你怎么炒股，也不解释因子公式——那是开发者手册和 docs/phase*-acceptance-report.md 的内容。",
    )

    doc.add_heading("2. 这是什么，不是什么", level=1)
    para(doc, "这是一台跑在你电脑上的投研辅助工具。它会：")
    bullets(
        doc,
        [
            "用已经训练好的排序模型，给全市场可投资股票打一个相对名次（不是预测涨跌幅）。",
            "根据你手动录入的持仓，计算浮盈亏、集中度、简化 VaR 和回撤。",
            "按固定优先级生成建议卡片：先看能不能交易，再看止损/集中度，再看模型信号，最后才谈再平衡。",
            "可选：把上述数字和新闻标题翻译成中文说明（需要你自己配置通义千问 API Key）。",
            "用 K 线图对照成本/止盈止损价；用历史建议复盘粗看过去提示方向对不对。",
        ],
    )
    para(doc, "它明确不做：")
    bullets(
        doc,
        [
            "连接券商、同步真实仓位、自动下单。",
            "保证收益、荐股必中、跟单致富。",
            "用大模型直接决定买还是卖（AI 只解释已经算好的结果）。",
            "高频、Level-2 抢跑、代客理财。",
        ],
    )

    doc.add_heading("2.1 排序模型在干什么（读这一段就够用）", level=2)
    para(
        doc,
        "系统给每只可投资股票打的是「这一天在全市场里相对靠前还是靠后」，不是「明天会涨百分之几」。"
        "训练时看的是：从某个月末起，未来最多约一个月里，价格是先碰到 +8%、先碰到 -8%，还是到期都没碰到。"
        "用来防自己骗自己的办法主要是：按日期切分训练/测试（同一天的股票不能一半拿去训练一半拿去考试）、"
        "严格用过去预测未来的滚动前推，以及回测时计入佣金、印花税、涨跌停买不进。"
        "所以首页和扫描里的分数只能当观察顺序，不能当下单指令。细节在开发者手册第 5.1 节。",
    )

    doc.add_heading("3. 第一次使用", level=1)
    doc.add_heading("3.1 启动（推荐：双击一键启动器）", level=2)
    para(
        doc,
        "项目根目录下的「研衡启动器.exe」——双击它，会自动打开一个黑色命令行窗口（用来显示后台日志，"
        "不用管它，不要关掉），几秒到几十秒后自动打开浏览器进入系统。首次启动稍慢，属正常现象。",
    )
    callout(
        doc,
        "为什么还留着黑色命令行窗口，不能更\u201c干净\u201d一点？",
        "本系统依赖较大的机器学习库（LightGBM、PyTorch、DuckDB），把它们全部塞进一个双击即用的"
        "exe 会让文件膨胀到几个GB且容易出现\u201c某个库在打包后的环境里找不到\u201d的问题。启动器只负责"
        "\u201c帮你按一次开始键\u201d，真正干活的还是项目自带、已经装好完整依赖的Python环境——这是权衡"
        "稳定性之后的选择，不是没打包干净。",
    )
    doc.add_heading("3.1.1 备用方式：命令行手动启动", level=2)
    para(doc, "如果启动器打不开（比如没有装 .venv 环境），可以在 PowerShell 中手动启动：")
    para(
        doc,
        'cd "e:\\妙妙工具\\炒股辅助\\stock-quant-system"\n'
        ".\\.venv\\Scripts\\python.exe -m streamlit run app.py",
        size=10,
        color=NAVY,
    )
    para(doc, "浏览器打开 http://localhost:8501 。左侧八个入口：首页（持仓驾驶舱）、持仓与建议、掘金扫描、风险仪表盘、AI 解释、行情图表、名词解释、历史建议复盘。")

    doc.add_heading("3.2 可选：配置 AI 解释", level=2)
    para(
        doc,
        "不配置也能用前三个功能。若要用「AI 解释」页：在项目根目录 .env 增加一行 DASHSCOPE_API_KEY=你的密钥"
        "（阿里云百炼：https://bailian.console.aliyun.com/ ）。没配时该页会黄条提示，结构化数据仍会显示。",
    )

    doc.add_heading("3.3 录入持仓", level=2)
    numbered(
        doc,
        [
            "打开「持仓与建议」。",
            "在搜索框输入 6 位数字或公司名称（如 紫金矿业 / 中信证券），多只匹配时在下拉框选定。不要带 .SH / .SZ。",
            "填股数、每股成本、建仓日期，提交。同一只股票可以分多批录入。",
            "「平仓」只把该批次标记为已结束，不会向券商发任何指令。",
        ],
    )
    callout(
        doc,
        "局限 · 持仓不是券商仓位",
        "系统不知道你在券商里实际买了多少、成交价里有没有滑点、现金有多少。"
        "集中度按「已录入持仓市值之和」当净值，现金不在模型里。漏录、错录都会让建议失真。",
    )

    doc.add_heading("4. 页面怎么用", level=1)
    doc.add_heading("4.1 持仓驾驶舱（首页）", level=2)
    para(doc, "看持仓只数、市值合计、浮动盈亏、当前生产冠军模型编号。单票占比超过 15% 会出警告。")
    para(
        doc,
        "首页新增「🎯 今日决策速览」：把最近一次在「持仓与建议」页生成的全部建议，按紧急程度"
        "（止损 > 减仓 > 止盈 > 再平衡 > 建仓 > 持有 > 观望）自动排序，只展示最需要关注的几条，"
        "并配一句大白话总结（比如「已经亏了9%，跌破止损线了，建议尽快考虑卖出止损」），不用逐页翻找。"
        "这个区块读的是历史记录，不会自动重新跑模型——要看最新结果，仍需去「持仓与建议」页点刷新。"
        "首页下半还有「按名称查代码」：输入紫金矿业、中信证券或 6 位代码，从本地股票池匹配，不联网。",
    )

    doc.add_heading("4.2 持仓与建议", level=2)
    para(
        doc,
        "点「生成/刷新建议卡片」会重新跑全市场打分，可能要几秒到几十秒。生成日期对应最近一个已入库的交易日，"
        "不一定是日历上的今天（周末或数据尚未更新时尤其如此）。",
    )
    make_table(
        doc,
        ["卡片动作", "大致含义", "你该怎么理解"],
        [
            ["watch", "观望 / 暂时无法操作", "停牌、涨跌停封死，或信号不够强。观望是正式动作，不是系统偷懒。"],
            ["open", "未持有、排名进入前 50", "候选观察，不是「立刻满仓」。"],
            ["hold", "继续持有", "没有触发止损/超限；不代表一定会涨。"],
            ["reduce", "建议降低仓位", "常见原因：单票超过净值 15%。"],
            ["stop_loss", "触及 -8% 止损参考线", "阈值来自模型训练标签，不是你的个性化止损。"],
            ["take_profit", "浮盈达 +8% 且模型排名已走弱", "涨了但排名仍很靠前时，不会单纯因为赚钱就让你卖。"],
            ["rebalance", "组合优化认为权重偏离过大", "只会出现在本来是 hold 的票上，不会覆盖止损。"],
        ],
    )
    para(doc, "每张卡片都有：理由（模型/因子/持仓）、风险提示、失效条件、免责声明。展开后请把四段都读完再决定是否去券商 App 操作。")
    para(
        doc,
        "卡片上现在还会直接换算出具体价格（现价 / 成本价 / 参考止损价 / 参考止盈价），"
        "不用自己心算「跌8%等于多少钱」——止损止盈价按你录入的加权平均成本价 × (1 ± 8%) 计算，"
        "和训练模型用的标签阈值完全一致。",
    )

    doc.add_heading("4.3 掘金扫描", level=2)
    para(
        doc,
        "模型分是相对排序，分高只表示在当天截面里更靠前。Top 表里的「行为金融提示」是冲突提示"
        "（浮盈筹码、龙虎榜过密、大小单背离），不会改排名，请自己交叉判断。点「❓ RankIC / 因子」可看大白话解释。",
    )
    para(
        doc,
        "2026-08-26 实测快照：交易日 2026-08-24，全市场候选 5340 只，可交易 5293 只。"
        "北交所小盘在历史上更容易排到前面，这是已知偏好，不是「这些公司更好」。滑块可调显示前 10～200 名。",
    )

    doc.add_heading("4.4 行情图表", level=2)
    para(
        doc,
        "输入 6 位代码或公司名称（如紫金矿业）看前复权 K 线 + 成交量。若这只股票已录入持仓，图上会画三条虚线："
        "持仓成本价、参考止损价、参考止盈价（与建议卡片同一套 ±8% 阈值）。未持有则只显示 K 线，并有说明文字。"
        "查不到数据时会出现友好提示，不会闪退。",
    )

    doc.add_heading("4.5 风险仪表盘", level=2)
    bullets(
        doc,
        [
            "集中度：水平柱状图（不是饼图），柱上直接标百分比，红色虚线是 15% 单票上限。",
            "波动率 / VaR：历史模拟法，用过去约 1 年真实日收益的经验分位数，不是银行那套参数法。",
            "相关性：两两日收益相关系数 > 0.7 时，「分散持仓」效果有限。",
            "回撤：按当前权重重放历史的假设性回撤，不是你真实调仓轨迹的回撤。点「❓ 最大回撤」可看解释。",
        ],
    )

    doc.add_heading("4.6 AI 解释", level=2)
    para(
        doc,
        "输入代码或公司名称后，左侧是系统算出来的分数、因子、行为信号和带发布时间的新闻；右侧是通义千问把这些内容翻译成中文。"
        "通义和 DeepSeek 都没有自己的实时日历，「今天」以本机日期为准，不是模型训练记忆里的年份。"
        "规则要求模型不得说「建议买入/卖出」，不得编造输入里没有的数字。新闻来自东方财富搜索，标题真实性由资讯源负责。"
        "没配 DASHSCOPE_API_KEY 时该页会黄条提示，结构化数据仍会显示。"
        "不要为了「模型还以为是 2024」去换 DeepSeek——那是训练截止日期，换模型解决不了。",
    )

    doc.add_heading("4.7 历史建议复盘", level=2)
    para(
        doc,
        "把过去每条建议和「发出当天收盘价 → 最新收盘价」的涨跌幅放在一起，标注「方向正确 / 方向不利」。"
        "可用「按建议类型筛选」勾选/清空类型。清空全部类型时应提示「没有符合当前筛选条件的记录」，而不是空白死页。"
        "如实说明：这只是简化统计（没考虑真实成交价、滑点、仓位大小），不是严格策略回测，更不是收益承诺。",
    )

    doc.add_heading("4.8 名词解释", level=2)
    para(
        doc,
        "集中解释 RankIC、VaR、置信度、PSI、因子等词。页面顶部可搜索；搜不到匹配项时列表为空，页面其余部分仍正常。"
        "各页关键指标旁的「❓」弹出的是同一份术语，不用来回切页面。",
    )

    doc.add_heading("5. 日常建议怎么配合使用", level=1)
    numbered(
        doc,
        [
            "先看风险仪表盘有没有集中度或高相关警告。",
            "再看持仓建议：止损、减仓优先于「我觉得还能涨」。",
            "未持有的票只把掘金 Top 当观察名单，结合冲突提示，不要按排名 1 到 10 依次满仓。",
            "去券商 App 下单时自行处理 T+1、涨跌停买不进/卖不出、最小交易单位。",
            "操作后回到本系统改持仓记录，避免下次建议还按旧仓位算。",
        ],
    )

    doc.add_heading("6. 数据不是实时行情", level=1)
    para(
        doc,
        "仓库里的行情、财务、停牌、涨跌停来自定期增量抓取（开发者跑 daily_pipeline）。"
        "盘中即时价格、集合竞价、你刚成交的那一笔，这里都没有。"
        "若建议日期停在几个交易日之前，先让开发者更新数据再生成卡片。",
    )

    doc.add_heading("7. 已知局限（请整节保留，不要跳过）", level=1)
    callout(
        doc,
        "局限是产品的一部分",
        "下面每一条都是验收时如实写下的边界，不是客套。按这些边界使用，系统是辅助；假装它们不存在，系统会变成误导。",
    )
    make_table(
        doc,
        ["局限", "对你意味着什么"],
        [
            ["持仓靠手录", "漏记、记错成本，建议和风险数字会一起错。"],
            ["模型是排序不是预测", "第 1 名不保证明天上涨；历史上头部组合也有大幅回撤（冠军回测最大回撤约 -34%）。"],
            ["止盈止损固定 ±8%", "来自训练标签，不是按你的风险承受力定制。"],
            ["VaR 是历史分位数", "极端新政、流动性枯竭不会提前出现在这条曲线里。"],
            ["组合优化很简化", "没有行业中性，没有换手上限；assumed_IC=0.03 是保守常数。"],
            ["行为信号是代理指标", "不是交易所官方筹码分布；当前数据权限买不到 cyq 接口。"],
            ["AI 可能仍会说滑", "即便 prompt 禁止下单用语，仍以结构化数字为准，不以 AI 段落为准。"],
            ["风格会变", "2026-08 监控显示动量因子 PSI 显著漂移，滚动夏普走弱；过时模型会给出过时排序。"],
            ["样本外超额曾经显著", "Phase 1 扣费后相对等权基准 p=0.0044，这是历史检验，不是未来保证。"],
        ],
    )

    doc.add_heading("8. 常见问题", level=1)
    para(doc, "Q：为什么建议减仓，明明在赚钱？", bold=True)
    para(doc, "A：单票占净值超过 15% 时，风控优先级高于「浮盈好看」。这是故意设计，不是 bug。")
    para(doc, "Q：AI 解释是灰的 / 黄条？", bold=True)
    para(doc, "A：没配 DASHSCOPE_API_KEY，或网络/额度失败。不影响持仓、扫描、风险三页。")
    para(doc, "Q：扫描结果全是北交所？", bold=True)
    para(doc, "A：已知小盘偏好。把扫描当研究输入，不要当购物清单。")
    para(doc, "Q：风险页怎么没有饼图了？", bold=True)
    para(doc, "A：持仓常常不止 4 只，饼图角度不好比。改成水平柱 + 15% 红线，更容易看出谁超限。")
    para(doc, "Q：历史复盘里股票名变成 nan？", bold=True)
    para(doc, "A：旧版本有过这个显示 bug，2026-08-26 已修。若仍出现请记下代码和日期。")
    para(doc, "Q：我能把建议转发给别人当荐股吗？", bold=True)
    para(doc, "A：不能。本工具按个人研究用途设计，没有适当性管理，也没有代客资质。")

    doc.add_heading("9. 第二轮人工可用性测试", level=1)
    para(
        doc,
        "本轮测试目的：确认字体够大、建议能看懂、新页面（图表/名词/复盘）和边界情况不会崩。"
        "逐步清单、通过标准、记录表见同目录《人工可用性测试指南》（docx 与 md 各一份）。"
        "测的时候请按指南勾选，不要凭记忆跳步；发现与指南「预期」不符的，记代码/操作路径/截图。",
    )

    doc.add_heading("10. 修订与责任", level=1)
    para(
        doc,
        "本说明书与 2026-08-26 的 Phase 4 + UI 两轮打磨状态对齐（版本 1.1）。"
        "软件升级后以仓库内 README、docs/ux-upgrade-notes.md 和最新验收报告为准。"
        "作者与开发辅助工具不对使用本系统产生的任何盈亏承担责任。",
    )
    return doc


def build_dev() -> Document:
    doc = _setup_doc("开发者使用说明书")
    cover(
        doc,
        "个人量化研究系统  ·  stock-quant-system",
        "开发者使用说明书",
        "架构 · 运维 · 模型生命周期 · 制度边界 · LLM · 开发 Skill",
        [
            "版本 1.1    编写日期 2026-08-26",
            "适用对象：维护本仓库的开发者（即系统所有者）",
            "强制规则：.cursor/rules/*.mdc 与 .cursor/skills/yanheng-dev-loop/",
        ],
    )

    doc.add_heading("1. 制度是否合理？有没有企业级标准？", level=1)
    para(
        doc,
        "结论先说：制度设计合理，对标的是企业级原则，落地的是「个人可重复级」工程，不是券商或对冲基金的生产平台。",
    )
    para(
        doc,
        "你定的两套规则——八闸门研发闭环、八荣八耻、冠军必须同协议打赢才能替换——对应的是公开、可引用的行业实践："
        "ThoughtWorks CD4ML（代码+数据+模型一起版本化，挑战者对照冠军晋升）、CRISP-DM（评估不可跳过）、"
        "量化研究里的 Purged K-Fold / Walk-Forward（见理论库 docs/03-量化方法/10-时序验证与标签工程.md）。"
        "Phase 3 已经按门禁执行：DL 的 Walk-Forward RankIC 均值 0.1112 显著高于 LightGBM 的 0.0608（p=0.0182），"
        "但同口径 Top-30 净收益 Sharpe 0.90 对 1.31，因此未替换冠军。这说明制度不是墙上的字。",
    )

    doc.add_heading("1.1 和企业级清单逐项对照", level=2)
    make_table(
        doc,
        ["企业级能力（CD4ML / MLOps）", "本仓库", "差距"],
        [
            ["实验与模型注册", "mlops/registry/<run_id>/ + champion.json + promotion_log.jsonl", "无 MLflow 服务器，无多用户实验板"],
            ["同协议晋升门禁", "已实现，DL 未晋升", "阈值 0.01 RankIC 带宽是工程拍定"],
            ["漂移监控", "PSI / WF 趋势 / 滚动夏普 / 预留生产 IC", "生产 IC 需 ≥20 个交易日，目前不足"],
            ["数据版本", "单文件 DuckDB", "无 DVC/lakeFS，无法一键复现「某日仓库」"],
            ["CI/CD", "无自动流水线；GitHub 远程 github.com/DomenRaven/yanheng", "没有测试门禁流水线，push 仍手工"],
            ["模型服务", "进程内 pickle + Streamlit", "无 REST、无金丝雀、无影子流量"],
            ["权限与审计", "单用户本机，advice_log 可追溯建议", "无 RBAC、无操作者身份"],
            ["公平性/合规持牌", "文案层免责 + Won't 清单", "不是持牌投顾，不能对客"],
            ["特征存储", "parquet 面板 + DuckDB", "不是 Feast 点时特征店"],
            ["风险模型", "历史法 VaR + 简化均值方差", "不是 Barra 多因子风险模型"],
        ],
    )
    para(
        doc,
        "成熟度（对照常见 0–4 / 1–5 的 MLOps 梯子）大约在「2：训练与监控可重复，发布仍手工」。"
        "原则已经像企业；自动化、数据血统、在线服务、组织流程还差三档。不要用「我们有 CD4ML 文档」对外声称企业级。",
    )
    para(doc, "若要靠近企业级，优先顺序是：git 远程与提交纪律 → 任务计划跑 daily_pipeline / 周末重训 → 带日期的仓库备份 → 等 prediction_log 够长再启用生产 IC。不要先上微服务。")

    doc.add_heading("1.2 八闸门使用时的注意", level=2)
    bullets(
        doc,
        [
            "闸门 4「测试与打分」禁止改评估口径来让数字好看。Phase 3 没有把 RankIC 单独当成晋升条件，就是在执行这一条。",
            "闸门 8「微调下一任务」已经发生过：Phase 4 对照产品边界才补上组合优化，并在验收报告里写明「原任务列表未单列」。",
            "八荣八耻第 7 条要求把不确定写出来。各 phase 验收报告的「已知局限」章节是制度的证据，删掉它们等于拆制度。",
            "Agent 执行时加载项目 Skill：.cursor/skills/yanheng-dev-loop/（SKILL.md + eight-gates.md + eight-honors.md）。规则原文禁止改写八荣八耻的八句话。",
        ],
    )

    doc.add_heading("2. 仓库地图", level=1)
    make_table(
        doc,
        ["目录", "职责", "摸库吗"],
        [
            ["ingestion/", "增量抓取，唯一写仓库的批量入口经 scripts/daily_pipeline.py", "写"],
            ["common/", "配置、DuckDB schema、Tushare 客户端、重试", "读写基础设施"],
            ["research/", "纯函数因子/标签/验证；panel.py 是 research 里唯一摸库模块", "仅 panel/eval/train 读"],
            ["advice/", "扫描 + 建议卡片，不改模型权重", "读 + 写 prediction_log/advice_log"],
            ["behavior/", "处置效应/拥挤/情绪代理，只做冲突提示", "读"],
            ["risk/", "持仓风险 + 组合优化", "读 positions 与行情"],
            ["llm/", "新闻抓取 + 翻译层", "不写库"],
            ["mlops/", "重训调度、漂移、registry", "读面板/日志"],
            ["pages/ + app.py", "Streamlit UI，禁止在页面里复制业务公式", "经上述模块"],
        ],
    )
    para(doc, "分层约定不要打破：A 股规则以 research/a_share_rules.py 和 limit_price / suspend_calendar 为准；财务用 EAV；DuckDB 单写者，禁止并行写主库。")

    doc.add_heading("3. 环境与密钥", level=1)
    numbered(
        doc,
        [
            "Python 3.12，虚拟环境 .venv/。新终端用 pip 前设置 $env:PIP_CONFIG_FILE 指向 .venv/pip.ini（系统 pip.ini 编码问题）。",
            "依赖见 requirements.txt。torch 必须用 cu124 索引安装，默认源会装成 CPU 版。",
            ".env（已 gitignore）：TUSHARE_TOKEN、DASHSCOPE_API_KEY。不要提交、不要写进文档示例的真实值。",
            "工作目录必须是 stock-quant-system/，相对路径 data/warehouse.duckdb、mlops/registry 才找得到。",
        ],
    )

    doc.add_heading("4. 日常运维命令", level=1)
    para(doc, "数据层唯一入口：")
    para(doc, "python -m scripts.daily_pipeline\npython -m scripts.daily_pipeline --only quotes,tushare_prices\npython -m scripts.daily_pipeline --skip fundamentals,income_statement", size=10, color=NAVY)
    para(doc, "研究与模型：")
    para(
        doc,
        "python -m research.panel --start 20160101\n"
        "python -m research.factor_eval\n"
        "python -m research.train_lightgbm\n"
        "python -m research.backtest\n"
        "python -m advice.scanner --top-n 50\n"
        "python -m mlops.retrain_schedule\n"
        "python -m mlops.drift_monitor\n"
        "python -m streamlit run app.py",
        size=10,
        color=NAVY,
    )
    para(
        doc,
        "建议用 Windows「任务计划程序」：每个交易日收盘后若干小时跑 daily_pipeline（财务类可每周）；"
        "每周末跑 retrain_schedule 与 drift_monitor。脚本本身只做一次原子操作，不内置常驻进程。",
    )
    callout(
        doc,
        "局限 · 数据源与网络",
        "新浪为主、东财为辅；本机代理会导致部分东财接口不稳定。Tushare 2000 积分档仍无 cyq_perf 等权限。"
        "单位必须按官方文档换算（vol 手→股、amount 千元→元），Phase 2 已修过一次 100× 事故，新增 Tushare 表时先查档再写。",
    )

    doc.add_heading("5. 模型生命周期（你必须守的门）", level=1)
    para(
        doc,
        "生产冠军指针：mlops/registry/champion.json，当前为 LightGBM run_id=20260826_123439。"
        "scanner 只加载冠军，不加载「文件夹名最新」。挑战者训练后必须用同一套 _date_level_splits、同一套三重障碍标签、"
        "同一套 research.backtest.run_portfolio_backtest 对比。只比 RankIC 不够。",
    )
    para(
        doc,
        "晋升记录追加 promotion_log.jsonl。结论可以是「不晋升」。删除挑战者文件等于销毁证据，禁止。",
    )
    para(
        doc,
        "标签与建议阈值必须同源：训练 label_config 为 profit_take=0.08、stop_loss=0.08、horizon 20 日；"
        "advice_engine 的止盈止损用同一对数。改一边必须改另一边并重训，否则逻辑分裂。",
    )

    doc.add_heading("5.1 训练在做什么、怎么防过拟合", level=2)
    para(
        doc,
        "生产模型是 LightGBM LambdaRank，不是回归预测涨跌幅。每个月末截面把股票排成相对名次；"
        "相关性标签来自三重障碍（先触 +8% / -8% / 最多 20 个交易日），映射到 {0,1,2}。"
        "特征是 13 个点时因子；同一天的全部股票必须分在同一折（按日期切分，不是按行随机切）。",
    )
    bullets(
        doc,
        [
            "标签与特征时间箭头：因子只用截面日及之前；标签看截面日之后最多 20 日路径。",
            "验证：Purged K-Fold + Embargo 防标签区间重叠泄漏；Walk-Forward expanding 才是部署口径（当前 OOS RankIC 均值 0.0608）。Purged K-Fold 会偏高，报告里两套都写、以 WF 为准。",
            "树正则：num_leaves=31、min_child_samples=50、learning_rate=0.05、n_estimators=200；特征 1%/99% 分位缩尾。",
            "回测另加涨跌停/停牌/成本，禁止只用样本内曲线验收。",
            "挑战者（行业图注意力）同协议对比后未晋升：RankIC 更好但 Top-30 净收益更差。",
            "局限：最终生产模型用全历史拟合，OOS 数字来自折内模型；单一 walk-forward 路径、北交所小盘偏好。详见 phase1 报告第 6 节。",
        ],
    )

    doc.add_heading("6. 建议引擎优先级（不要改顺序「优化体验」）", level=1)
    numbered(
        doc,
        [
            "合规与可交易性（停牌 / 涨停买不进 / 跌停卖不出）→ watch",
            "生存风控（-8% 止损、单票 >15%）→ stop_loss / reduce",
            "账户结构",
            "信号质量（+8% 且排名走弱才 take_profit）",
            "收益增强：apply_rebalance_overlay 只能把 hold 升级为 rebalance",
        ],
    )
    para(doc, "单测约定：reduce 不得被优化器覆盖。组合优化失败要吞掉异常，不能阻断主建议。")

    doc.add_heading("7. 大模型：DeepSeek 还是通义千问", level=1)
    para(
        doc,
        "本系统的 LLM 是翻译层，不是决策层。选型标准因此不是「谁更像 GPT」，而是：指令遵守（禁止买卖用语）、"
        "不编造输入没有的数字、中文可读、接口稳定、改动能不能局部化。",
    )
    make_table(
        doc,
        ["维度", "DeepSeek", "通义千问（DashScope，当前默认）"],
        [
            ["接口", "OpenAI 兼容，api.deepseek.com", "OpenAI 兼容，dashscope.aliyuncs.com（当前默认）"],
            ["代码改动", "llm/explain_assistant.py 已接好", "config.yaml 的 llm.provider=dashscope + DASHSCOPE_API_KEY"],
            ["任务匹配", "遵守系统 prompt 足够；中文偏书面", "说明文案通常更顺；当前生产用 qwen-plus"],
            ["成本", "个人点选", "按阿里云百炼计费"],
            ["运维", "独立账号", "阿里云企业合同、发票更完整"],
            ["本项目现状", "曾为默认，已切换走", "2026-08-26 起为默认解释层"],
        ],
    )
    para(
        doc,
        "当前默认是通义千问（DashScope / qwen-plus）。解释质量的上限是输入质量（模型分数、因子、带时间戳的新闻），"
        "换一个同级中文模型几乎不会改变投资决策。"
        "双方都没有实时日历：系统会把本机日期写进提示词。换 DeepSeek 不能让模型「知道今天」。"
        "不要默认打开通义 extra_body.enable_search（兼容接口不返回搜索来源）。"
        "切换提供商：改 config.yaml 的 llm.provider 与 api_key_env，不要复制一套 explain_assistant。"
        "价格与限流以官网当天页面为准。",
    )
    callout(
        doc,
        "局限 · LLM",
        "Key 必须使用者自己申请。未配置时必须优雅降级。禁止让 LLM 输出 action 或 pred_score。"
        "新闻抓取绕开了 akshare stock_news_em 在本环境 pandas pyarrow 后端上的正则 bug，不要「为了复用」改回那个函数。",
    )

    doc.add_heading("8. 如何扩展（保持小步）", level=1)
    bullets(
        doc,
        [
            "新因子：写在 research/factors.py（纯函数），在 panel.py 接入 point-in-time，factor_eval 出 IC，再决定是否进训练特征。禁止直接在 Streamlit 里算因子。",
            "新数据源：先小样本打真实请求核对字段，再进 ingestion/，最后挂上 daily_pipeline 步骤表。",
            "新建议动作：先改 advice_engine 优先级文档字符串和单测，再改 UI 表情映射。",
            "一次只做一个 Phase 量级的改动。跨分层（例如让 llm 写 advice_log 的 action）先写原因。",
        ],
    )

    doc.add_heading("9. 已知局限与刻意不做（开发视角）", level=1)
    make_table(
        doc,
        ["项目", "状态", "不要假装已解决"],
        [
            ["冠军回测净 CAGR / Sharpe", "54.7% / 1.31，相对等权超额 p=0.0044", "单条 walk-forward、84 个月，过拟合风险仍在"],
            ["DL 挑战者", "RankIC 更好，Top-N 更差，未上线", "未做分位加权 loss，原因仍是假设"],
            ["ST 历史标记", "universe.is_st 是当前快照", "回测无法完美还原历史 ST"],
            ["行业/风格中性", "组合优化未做", "产品 Should 里「中性化」仍部分缺口"],
            ["券商 API", "Won't", "不要为了「方便」去接未持牌代下单"],
            ["git", "已推送 github.com/DomenRaven/yanheng（仓库级 user.name）", "无 CI；勿提交 .env / duckdb / exe"],
        ],
    )

    doc.add_heading("9a. 一键启动 exe 打包（UI优化里新增）", level=1)
    para(
        doc,
        "打包脚本：python scripts/build_exe.py（需先 pip install pyinstaller，仅打包时需要，不是"
        "运行时依赖）。产物是项目根目录下的「研衡启动器.exe」，约几MB。",
    )
    callout(
        doc,
        "工程决策：为什么不把整个应用连同 torch/lightgbm/duckdb 一起冻结成单文件exe",
        "PyInstaller把这类含C扩展、体积巨大的库整体冻结进单文件，社区里公认容易踩两类坑："
        "一是产物体积轻松膨胀到数GB且难以精简；二是Streamlit的静态资源、torch的动态库在frozen模式下"
        "经常需要额外写hook才能找到，出问题后排查成本很高。因此选择更稳妥的方案：exe只承担"
        "\u201clauncher.py\u201d这一段轻量逻辑（找项目目录→调用.venv里的python→跑streamlit→开浏览器），"
        "真正的重依赖仍然用开发机上已经装好的.venv，效果同样是双击即用，但不需要为了"
        "\u201c看起来更极致\u201d去冒生产稳定性的风险。这是一个明确的工程取舍，不是打包能力不足。",
    )
    para(
        doc,
        "launcher.py 逻辑：定位exe所在目录 → 检查 .venv/Scripts/python.exe 和 app.py 是否存在 → "
        "若8501端口已被占用直接打开浏览器（避免重复启动）→ 否则用 .venv 的 python 启动"
        "`streamlit run app.py --server.headless true` → 轮询端口最多60秒 → 打开默认浏览器。"
        "换新机器分发时，仍需先完整走一遍第3节的环境搭建（python -m venv .venv + pip install -r "
        "requirements.txt），exe本身不携带依赖，这是设计如此，不是遗漏。",
    )
    para(doc, "修改 launcher.py 后需要重新执行 python scripts/build_exe.py 才会反映到exe里；该脚本会自动清理 PyInstaller 产生的中间文件（_dist_tmp/_build_tmp），不需要手动收尾。")

    doc.add_heading("10. 验收报告与再生本说明书", level=1)
    para(
        doc,
        "权威数字以 stock-quant-system/docs/phase0~4-acceptance-report.md 为准，不要从聊天记录抄。"
        "修改手册后可运行：python -m scripts.build_manuals ，输出 docs/manuals/ 下用户手册、开发者手册、人工可用性测试指南三份 docx。"
        "Agent 开发约束见仓库根目录 .cursor/skills/yanheng-dev-loop/ 。",
    )
    para(
        doc,
        "本说明书与软件同样不构成投资建议。开发者改了评估口径却不改报告，属于违反 quant-dev-loop 闸门 4。",
    )
    return doc


def build_usability() -> Document:
    doc = _setup_doc("人工可用性测试指南")
    cover(
        doc,
        "个人量化研究系统  ·  研衡 YanHeng",
        "人工可用性测试指南",
        "第二轮（UI 理论复核后）· 请按序勾选，不要跳步",
        [
            "版本 1.0    编写日期 2026-08-26",
            "适用对象：系统所有者本人（非开发者视角的「能不能用来辅助决策」）",
            "配套说明书：同目录《用户使用说明书》版本 1.1",
        ],
    )

    doc.add_heading("0. 这一轮测什么、不测什么", level=1)
    para(
        doc,
        "测：打开、看懂、点完一条「录入持仓 → 生成建议 → 对照图表/风险/复盘」的决策路径；"
        "字体是否够大；边界输入会不会崩。",
    )
    para(
        doc,
        "不测：模型重训、全市场回测统计显著性、数据管道回填。那些已在 Phase 1～3 验收报告里用真实数据打过分。"
        "本指南是人工可行性/可用性，不是量化打分。",
    )
    callout(
        doc,
        "通过标准（本轮）",
        "8 个页面都能打开且无红色 Python 报错；主路径（持仓→建议→首页速览）能走通；"
        "指南第 3 节边界用例全部是「友好提示」而不是崩溃；你能不查代码就说出「止损价是怎么来的」。"
        "主观项（字是否够大、卡片是否好懂）请在记录表打分 1–5，低于 3 的请写下原因。",
    )

    doc.add_heading("1. 测试前准备", level=1)
    numbered(
        doc,
        [
            "确认仓库目录存在 .venv 和 data/warehouse.duckdb（没有数据库则扫描/图表会空，属环境问题不是 UI 问题）。",
            "推荐：双击项目根目录「研衡启动器.exe」。备用：在 stock-quant-system 下运行 .venv\\Scripts\\python.exe -m streamlit run app.py",
            "浏览器打开 http://localhost:8501 。不要同时开两个启动器抢 8501 端口。",
            "准备两只真实持仓做对照：建议一只普通（如 000001）、一只你仓位较重的（测 15% 集中度）。代码填 6 位数字、不要带 .SH。",
            "本文件打印或另开窗口，边测边勾。不通过的记下：页面名 / 你点了什么 / 看到什么 / 是否有红色报错。",
        ],
    )

    doc.add_heading("2. 主路径（必须全过）", level=1)
    make_table(
        doc,
        ["步骤", "操作", "预期（通过）", "不通过长什么样"],
        [
            ["2.1 首页", "打开默认页", "标题「研衡」；免责声明；四格指标；有或没有「今日决策速览」都正常", "整页红字 traceback"],
            ["2.2 录入", "持仓与建议：填代码+股数>0+成本>0，提交", "出现成功提示，下方列表出现该批次", "无提示或闪退"],
            ["2.3 空表单", "不选股票再提交；或股数填 0", "红字要求先选定股票且股数/成本大于0", "Python 异常"],
            ["2.4 生成建议", "点「生成/刷新建议卡片」，等几十秒", "持仓票出现止损/减仓/止盈/持有等卡片；有大白话+具体价格", "一直转圈后红字；或只有百分比没有价格"],
            ["2.5 回首页", "回到持仓驾驶舱", "今日决策速览出现刚才的卡片，风控类排在前面", "仍提示还没生成过"],
            ["2.6 问号帮助", "回首页找「持仓明细」右侧问号；无持仓时标题仍在。也可在「持仓与建议」的当前持仓旁点问号", "弹出「集中度」解释", "找不到标题（应在首页下半段，不在持仓与建议页中部）"],
        ],
    )

    doc.add_heading("3. 逐页检查（含边界）", level=1)
    doc.add_heading("3.1 掘金扫描", level=2)
    bullets(
        doc,
        [
            "滑块拉到 10，点运行；再拉到 200 再运行一次。两次都应出表，显示快照交易日和可交易只数。",
            "表头应是中文（排名/代码/名称），不是纯英文列名。",
            "点问号 RankIC、问号因子，应弹出大白话。",
        ],
    )
    doc.add_heading("3.2 行情图表", level=2)
    bullets(
        doc,
        [
            "输入已持有代码：应有 K 线，且能看到成本/止损/止盈虚线（若该票在持仓里）。",
            "输入 999999 回车：提示「没有查到这只股票的行情数据」，不崩溃。",
            "字号：坐标轴和指标数字应明显大于旧版 Streamlit 默认（主观 ≥3 分算过）。",
        ],
    )
    doc.add_heading("3.3 风险仪表盘", level=2)
    bullets(
        doc,
        [
            "集中度必须是水平柱状图，不是饼图；应能看到「15%风控上限」虚线。",
            "若单票超过 15%，顶部应有红色警告。",
            "问号 VaR、问号最大回撤可弹出解释。",
        ],
    )
    doc.add_heading("3.4 AI 解释", level=2)
    bullets(
        doc,
        [
            "输入一只真实代码（如 000001）：左侧结构化数据要有；右侧若已配千问 Key，应出中文说明且不应出现「建议买入/立刻满仓」这类指令口吻。",
            "左侧「系统日历日期」应是本机当天，新闻行应带发布时间；右侧不应把「今天」说成 2024（那是训练记忆，不是断网，也不用换 DeepSeek）。",
            "输入 abcdef 或 000000：不应崩溃；可以空数据或模型说明代码无效。",
            "未配 Key：黄条提示，左侧数字仍在。",
        ],
    )
    doc.add_heading("3.5 名词解释", level=2)
    bullets(
        doc,
        [
            "无搜索词时有一列术语卡片。",
            "搜索 xyz123：结果为空，页脚说明还在，不报错。",
            "搜索 VaR 或回撤：能搜到对应条。",
        ],
    )
    doc.add_heading("3.6 历史建议复盘", level=2)
    bullets(
        doc,
        [
            "生成过建议后本页应有记录。股票名不应出现字面 nan。",
            "筛选器清空全部类型：出现「没有符合当前筛选条件的记录」，不是一片死白。",
            "标题旁问号「置信度」可弹出解释。",
        ],
    )

    doc.add_heading("4. 观感与决策是否好懂（主观，请打分）", level=1)
    make_table(
        doc,
        ["问题", "1 很差 … 5 很好", "备注（低于 3 必填）"],
        [
            ["正文和按钮字是否够大、不用凑近看", "", ""],
            ["建议卡片能否 10 秒内看懂「要我干什么」", "", ""],
            ["止损/止盈是否看到了人民币价格而不只是 8%", "", ""],
            ["首页速览是否让你少翻页", "", ""],
            ["减仓/止盈等色块文字是否够清楚（对比度）", "", ""],
        ],
    )

    doc.add_heading("5. 已知现象（请勿当成新 bug）", level=1)
    make_table(
        doc,
        ["现象", "原因", "怎么处理"],
        [
            ["建议日期不是日历上的今天", "用最近已入库交易日", "周末或数据未更新时正常"],
            ["扫描前列很多北交所", "已知小盘+价值因子偏好", "当观察名单，不要当购物清单"],
            ["生成建议要等几十秒", "要跑全市场打分", "等转圈结束；不要连点多次（库内已按股票+日期去重）"],
            ["启动器还会弹出黑色命令行", "设计如此，日志窗口", "不要关；关了应用会停"],
            ["exe 换电脑双击没用", "exe 不含 Python 依赖", "先搭 .venv 再用启动器"],
            ["控制台提到 use_container_width", "Streamlit 弃用警告", "功能仍正常，可忽略"],
        ],
    )

    doc.add_heading("6. 记录表（测完请填）", level=1)
    make_table(
        doc,
        ["项", "结果（通过 / 不通过 / 跳过）", "说明"],
        [
            ["2 主路径 2.1–2.6", "", ""],
            ["3.1 掘金扫描", "", ""],
            ["3.2 行情图表", "", ""],
            ["3.3 风险仪表盘", "", ""],
            ["3.4 AI 解释", "", ""],
            ["3.5 名词解释", "", ""],
            ["3.6 历史建议复盘", "", ""],
            ["第 4 节主观均不低于 3", "", ""],
            ["本轮总评", "", "通过 / 有缺陷但可用 / 不通过"],
        ],
    )
    para(
        doc,
        "测完后：把「不通过」项发给后续模型时，请附页面名、操作步骤、报错原文。不要只说「有问题」。"
        "本指南不构成投资建议；测试持仓请用你愿意暴露给本机数据库的真实或模拟仓位。",
    )
    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    user_path = OUT_DIR / "用户使用说明书.docx"
    dev_path = OUT_DIR / "开发者使用说明书.docx"
    test_path = OUT_DIR / "人工可用性测试指南.docx"
    build_user().save(user_path)
    build_dev().save(dev_path)
    build_usability().save(test_path)
    print(user_path)
    print(dev_path)
    print(test_path)


if __name__ == "__main__":
    main()
